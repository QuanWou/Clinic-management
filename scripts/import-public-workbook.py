import datetime
import hashlib
import json
import os
import subprocess
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = Path(sys.argv[1])
CONFIG_PATH = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / ".runtime" / "main" / "config.json"
CONFIG = json.loads(CONFIG_PATH.read_text(encoding="utf-8-sig"))
DATABASE_NAME = CONFIG["databaseName"]

BASE = {
    "auth": "http://127.0.0.1:8093/modules/auth",
    "identity": "http://127.0.0.1:8093/modules/identity",
    "clinic": "http://127.0.0.1:8094/modules/clinic",
    "doctor": "http://127.0.0.1:8094/modules/doctor",
    "search": "http://127.0.0.1:8094/modules/search",
    "catalog": "http://127.0.0.1:8103/modules/catalog",
}


def api(service, path, method="GET", token=None, body=None, key=None):
    data = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")
    headers = {"Accept": "application/json"}
    if data is not None:
        headers["Content-Type"] = "application/json; charset=utf-8"
    if token:
        headers["Authorization"] = "Bearer " + token
    if key:
        headers["Idempotency-Key"] = key
    request = urllib.request.Request(BASE[service] + path, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request, timeout=25) as response:
            raw = response.read().decode("utf-8-sig")
            return json.loads(raw) if raw else None
    except urllib.error.HTTPError as error:
        raw = error.read().decode("utf-8-sig", errors="replace")
        try:
            detail = json.loads(raw)
        except Exception:
            detail = raw
        raise RuntimeError(f"{service} {method} {path}: HTTP {error.code} {detail}") from error


def register_and_login(email, full_name):
    password = CONFIG["testPassword"]
    try:
        api("auth", "/api/auth/register", "POST", body={"email": email, "password": password, "fullName": full_name})
    except RuntimeError as error:
        if "HTTP 409" not in str(error):
            raise
    result = api("auth", "/api/auth/login", "POST", body={"email": email, "password": password})
    payload = result.get("data", result)
    return payload["userId"], payload["accessToken"]


def psql(sql, capture=False):
    environment = os.environ.copy()
    environment["PGPASSWORD"] = CONFIG["databasePassword"]
    command = [
        str(Path(CONFIG["postgresBin"]) / "psql.exe"), "-X", "-w", "-h", CONFIG["databaseHost"],
        "-p", str(CONFIG["databasePort"]), "-U", CONFIG["databaseUser"], "-d", DATABASE_NAME,
        "-v", "ON_ERROR_STOP=1",
    ]
    if capture:
        command += ["-A", "-t", "-c", sql]
    result = subprocess.run(
        command,
        input=None if capture else sql.encode("utf-8"),
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        env=environment,
        check=False,
    )
    if result.returncode:
        raise RuntimeError(result.stderr.decode("utf-8", errors="replace"))
    return result.stdout.decode("utf-8", errors="replace").strip()


def sql_literal(value):
    return "'" + str(value).replace("'", "''") + "'"


def workbook_records(workbook, sheet):
    populated = [list(row) for row in workbook[sheet].iter_rows(values_only=True) if any(value is not None for value in row)]
    headers = populated[0]
    return [dict(zip(headers, row)) for row in populated[1:]]


def split_list(value):
    return [part.strip() for part in str(value or "").split("|") if part and part.strip()]


def main():
    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    clinic_rows = [list(row) for row in workbook["Phòng khám"].iter_rows(values_only=True) if any(value is not None for value in row)]
    clinic_fields = {str(row[0]): row[1] for row in clinic_rows[1:]}
    specialties = workbook_records(workbook, "Chuyên khoa")
    doctors = workbook_records(workbook, "Bác sĩ")
    services = workbook_records(workbook, "Dịch vụ")

    owner_id, owner_token = register_and_login("owner.publictest@clinic-v2.local", "Quản trị Clinic V2")
    reviewer_id, reviewer_token = register_and_login("reviewer.publictest@clinic-v2.local", "Kiểm duyệt Clinic V2")
    psql(f"INSERT INTO iam.platform_operators(user_id) VALUES({sql_literal(reviewer_id)}) ON CONFLICT(user_id) DO NOTHING;")

    slug = "phong-kham-da-khoa-clinic-v2"
    clinic_id = psql(f"select id from clinic.clinics where slug={sql_literal(slug)} limit 1", capture=True)
    license_input = {
        "licenseNumber": "PUBLIC-TEST-DATA",
        "issuingAuthority": "LOCAL-TEST",
        "scopeSummary": "Synthetic public website fixture only",
        "evidenceRef": "local-test/public-workbook",
        "validUntil": (datetime.date.today() + datetime.timedelta(days=365)).isoformat(),
    }
    draft = {
        "name": clinic_fields["Tên phòng khám"],
        "slug": slug,
        "publicDescription": clinic_fields.get("Giới thiệu ngắn"),
        "contactName": "Liên hệ nội bộ dữ liệu kiểm thử",
        "contactEmail": "owner.publictest@clinic-v2.local",
        "contactPhone": "0000000000",
        "license": license_input,
    }
    if clinic_id:
        clinic = api("clinic", f"/api/clinics/{clinic_id}", token=owner_token)
        if clinic["reviewStatus"] in ("DRAFT", "NEEDS_CHANGES", "APPROVED") and clinic["publicationStatus"] == "UNPUBLISHED":
            clinic = api("clinic", f"/api/clinics/{clinic_id}", "PUT", owner_token, {**draft, "expectedVersion": clinic["version"]})
    else:
        clinic = api("clinic", "/api/clinics", "POST", owner_token, draft, "workbook-clinic-create")
        clinic_id = clinic["id"]

    api("clinic", f"/api/clinics/{clinic_id}/owner-membership", "POST", owner_token, {}, "workbook-owner-membership")
    clinic = api("clinic", f"/api/clinics/{clinic_id}", token=owner_token)
    if clinic.get("branches"):
        branch_id = clinic["branches"][0]["id"]
    else:
        clinic = api(
            "clinic", f"/api/clinics/{clinic_id}/branches", "POST", owner_token,
            {"name": clinic_fields["Tên phòng khám"], "address": clinic_fields["Địa chỉ"], "openingHours": "Theo lịch hẹn", "active": True, "expectedVersion": clinic["version"]},
            "workbook-branch-create",
        )
        branch_id = clinic["branches"][0]["id"]

    clinic = api("clinic", f"/api/clinics/{clinic_id}", token=owner_token)
    if clinic["reviewStatus"] == "DRAFT":
        clinic = api("clinic", f"/api/clinics/{clinic_id}/submit", "POST", owner_token, {}, "workbook-submit")
    if clinic["reviewStatus"] == "SUBMITTED":
        clinic = api("clinic", f"/api/platform/clinics/{clinic_id}/approve", "POST", reviewer_token, {"reason": "LOCAL TEST ONLY: supplied workbook fixture", "evidenceVerified": True}, "workbook-approve")
    if clinic["publicationStatus"] == "UNPUBLISHED":
        clinic = api("clinic", f"/api/platform/clinics/{clinic_id}/publish", "POST", reviewer_token, {"reason": "LOCAL TEST ONLY: publish supplied public website fixture"}, "workbook-publish")

    memberships = api("identity", f"/api/clinics/{clinic_id}/memberships", token=owner_token)
    affiliations = api("doctor", f"/api/clinics/{clinic_id}/branches/{branch_id}/doctor-affiliations", token=owner_token)
    public_doctors = []
    # The supplied workbook deliberately leaves clinic opening hours unconfigured,
    # but the local booking acceptance fixture still needs real Doctor source
    # schedules. Keep this synthetic availability in Doctor Service instead of
    # inventing availability in the public UI or Appointment Service.
    schedule_from = (datetime.date.today() - datetime.timedelta(days=1)).isoformat()
    for row in doctors:
        code = str(row["Mã bác sĩ"])
        user_id, doctor_token = register_and_login(code.lower() + "@clinic-v2.local", str(row["Họ tên"]))
        membership = next((item for item in memberships if item["userId"] == user_id and item["role"] == "DOCTOR" and item["status"] != "REVOKED"), None)
        if membership is None:
            membership = api(
                "identity", f"/api/clinics/{clinic_id}/memberships", "POST", owner_token,
                {"userId": user_id, "role": "DOCTOR", "allBranches": False, "reason": "Import hồ sơ bác sĩ từ bộ dữ liệu public test"},
                "membership-" + code,
            )
            memberships.append(membership)
        if membership["status"] == "INVITED":
            membership = api("identity", f"/api/memberships/{membership['id']}/activate", "POST", doctor_token, {}, "activate-" + code)
        if branch_id not in membership.get("branchIds", []):
            api(
                "identity", f"/api/clinics/{clinic_id}/memberships/{membership['id']}/branches", "POST", owner_token,
                {"branchId": branch_id, "reason": "Phân công tại cơ sở public test"}, "branch-" + code,
            )
        affiliation = next((item for item in affiliations if item["userId"] == user_id and item["active"]), None)
        if affiliation is None:
            affiliation = api(
                "doctor", f"/api/clinics/{clinic_id}/branches/{branch_id}/doctor-affiliations", "POST", owner_token,
                {
                    "userId": user_id, "displayName": str(row["Họ tên"]), "registrationCode": code,
                    "specialtyCode": str(row["Specialty slug"]), "specialtyName": str(row["Chuyên khoa"]),
                    "professionalTitle": str(row["Danh xưng"]),
                    "effectiveFrom": (datetime.date.today() - datetime.timedelta(days=1)).isoformat(),
                    "effectiveUntil": None, "publicVisible": True,
                },
                "affiliation-" + code,
            )
            affiliations.append(affiliation)
        schedule_path = f"/api/clinics/{clinic_id}/branches/{branch_id}/doctor-affiliations/{affiliation['id']}/schedules"
        schedule_view = api("doctor", schedule_path, token=owner_token)
        schedules = schedule_view.get("schedules", [])
        for day_of_week in range(1, 7):
            exists = any(
                item["dayOfWeek"] == day_of_week
                and str(item["startTime"])[:5] == "08:00"
                and str(item["endTime"])[:5] == "17:00"
                and item.get("active", True)
                for item in schedules
            )
            if not exists:
                created = api(
                    "doctor", schedule_path, "POST", owner_token,
                    {
                        "dayOfWeek": day_of_week, "startTime": "08:00", "endTime": "17:00",
                        "effectiveFrom": schedule_from, "effectiveUntil": None,
                        "timezone": "Asia/Ho_Chi_Minh", "active": True,
                    },
                    f"public-test-schedule-{code}-{day_of_week}",
                )
                schedules.append(created)
        public_doctors.append({
            "sourceDoctorId": str(row["Doctor ID"]), "doctorId": affiliation["practitionerId"], "code": code,
            "displayName": str(row["Họ tên"]), "title": str(row["Danh xưng"]), "specialtyName": str(row["Chuyên khoa"]),
            "specialtySlug": str(row["Specialty slug"]), "yearsExperience": int(row["Số năm kinh nghiệm"]),
            "headline": str(row["Dòng chức danh"]), "summary": str(row["Giới thiệu ngắn"]),
            "expertise": split_list(row["Chuyên môn nổi bật"]), "consultationAreas": split_list(row["Khám và tư vấn"]),
            "education": split_list(row["Quá trình đào tạo"]), "experience": split_list(row["Quá trình công tác"]), "cta": str(row["CTA"]),
        })

    offering_rows = api("catalog", f"/api/clinics/{clinic_id}/offerings", token=owner_token)
    assignment_rows = api("catalog", f"/api/clinics/{clinic_id}/branches/{branch_id}/offerings", token=owner_token)
    public_services = []
    effective = (datetime.datetime.now(datetime.timezone.utc) - datetime.timedelta(days=1)).isoformat().replace("+00:00", "Z")
    for row in services:
        code = str(row["Mã dịch vụ"])
        offering = next((item for item in offering_rows if item["code"] == code), None)
        if offering is None:
            offering = api(
                "catalog", f"/api/clinics/{clinic_id}/offerings", "POST", owner_token,
                {"code": code, "name": str(row["Tên dịch vụ"]), "description": str(row["Mô tả"]), "specialtyCode": str(row["Specialty slug"]), "active": True},
                "offering-" + code,
            )
            offering_rows.append(offering)
        assignment = next((item for item in assignment_rows if item["offering"]["id"] == offering["id"]), None)
        if assignment is None:
            assignment = api(
                "catalog", f"/api/clinics/{clinic_id}/branches/{branch_id}/offerings", "POST", owner_token,
                {"offeringId": offering["id"], "durationMinutes": 30, "active": True, "publicVisible": True}, "assignment-" + code,
            )
            assignment_rows.append(assignment)
        prices = api("catalog", f"/api/clinics/{clinic_id}/branches/{branch_id}/offerings/{offering['id']}/price-versions", token=owner_token)
        if not prices:
            api(
                "catalog", f"/api/clinics/{clinic_id}/branches/{branch_id}/price-versions", "POST", owner_token,
                {"offeringId": offering["id"], "amountVnd": 0, "effectiveFrom": effective}, "price-" + code,
            )
        public_services.append({
            "offeringId": offering["id"], "code": code, "name": str(row["Tên dịch vụ"]),
            "specialtyName": str(row["Chuyên khoa"]), "specialtySlug": str(row["Specialty slug"]),
            "description": str(row["Mô tả"]), "suitableFor": str(row["Phù hợp với"]),
            "preparation": str(row["Chuẩn bị trước khi đến"]), "cta": str(row["CTA"]),
        })

    content = {
        "clinic": {
            "heroMessage": clinic_fields.get("Thông điệp hero"), "shortIntroduction": clinic_fields.get("Giới thiệu ngắn"),
            "detailedIntroduction": clinic_fields.get("Giới thiệu chi tiết"), "facilities": clinic_fields.get("Cơ sở vật chất"),
            "careProcess": clinic_fields.get("Quy trình khám"),
        },
        "specialties": [{
            "specialtyId": str(row["Specialty ID"]), "databaseName": str(row["Tên DB"]), "displayName": str(row["Tên hiển thị"]),
            "slug": str(row["Slug"]), "shortDescription": str(row["Mô tả ngắn"]), "description": str(row["Giới thiệu chi tiết"]),
            "commonConditions": split_list(row["Bệnh/vấn đề thường gặp"]), "keyExpertise": split_list(row["Chuyên môn chính"]),
            "doctorCount": int(row["Số bác sĩ"]), "cta": str(row["CTA"]),
        } for row in specialties],
        "doctors": public_doctors,
        "services": public_services,
    }
    source_hash = hashlib.sha256(WORKBOOK.read_bytes()).hexdigest()
    payload = json.dumps(content, ensure_ascii=False, separators=(",", ":"))
    psql(
        "INSERT INTO search_v2.public_web_content(clinic_id,content,source_label,source_sha256) VALUES(" +
        sql_literal(clinic_id) + "," + sql_literal(payload) + "::jsonb," + sql_literal(WORKBOOK.name) + "," + sql_literal(source_hash) + ") " +
        "ON CONFLICT(clinic_id) DO UPDATE SET content=excluded.content,source_label=excluded.source_label,source_sha256=excluded.source_sha256,imported_at=now();"
    )

    for _ in range(45):
        try:
            projected_doctors = api("search", f"/api/public/clinics/{clinic_id}/doctors")
            projected_services = api("search", f"/api/public/clinics/{clinic_id}/offerings")
            if len({item["doctorId"] for item in projected_doctors}) == len(doctors) and len({item["offeringId"] for item in projected_services}) == len(services):
                break
        except Exception:
            pass
        time.sleep(1)
    else:
        raise RuntimeError("Timed out waiting for public search projections")

    CONFIG["clinic"] = clinic_id
    CONFIG["branch"] = branch_id
    CONFIG_PATH.write_text(json.dumps(CONFIG, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"clinicId": clinic_id, "branchId": branch_id, "specialties": len(specialties), "doctors": len(doctors), "services": len(services), "sha256": source_hash}, ensure_ascii=True))


if __name__ == "__main__":
    main()
